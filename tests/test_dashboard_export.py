"""Pruebas automatizadas para exportación de reportes Excel/CSV y entregables BI."""

import io
import json
import zipfile
from fastapi.testclient import TestClient
import openpyxl
import pandas as pd
import pytest
from sqlalchemy import create_engine

import cli
from dashboard import app as dash_app
from dashboard.auth import engine as auth_engine, users as auth_users
from licensing.client import License, LicenseManager


def _licensed_manager(modules=None):
    if modules is None:
        modules = {"dashboard": True, "derived": True, "bi": True}
    lic = License(
        customer_id="cust_test",
        license_id="LIC-TEST-001",
        status="active",
        valid_from=0,
        valid_until=9999999999,
        offline_until=9999999999,
        modules=modules,
        limits={},
        issued_at=1000,
    )
    mgr = LicenseManager(config={})
    mgr._cached = {"token": "dummy", "claims": lic.__dict__, "fetched_at": 1000}
    mgr.get_license = lambda: lic
    return mgr


@pytest.fixture
def export_env(tmp_path, monkeypatch):
    db_path = tmp_path / "control.db"
    cfg_path = tmp_path / "config.json"
    cfg = {
        "environment": "qa",
        "source": {"type": "csv", "config": {"directory": str(tmp_path)}},
        "database": {"dialect": "sqlite", "database": str(db_path)},
        "tables": [{"source": "MARA", "target": "Mara_Data", "keys": ["MATNR"]}],
        "fields": {},
        "derived": [{
            "name": "inv_bodega",
            "tab": "Inventario Bodega",
            "from": "Mara_Data",
            "keys": ["MATNR"],
            "columns": [
                {"as": "ID", "source": "MATNR"},
                {"as": "Descripcion", "source": "MAKTX"},
                {"as": "StockLibre", "source": "LABST"},
                {"as": "stock_minimo", "source": "MINBE"},
                {"as": "ValorTotal", "source": "SALK3"},
                {"as": "Centro", "source": "WERKS"},
                {"as": "Almacen", "source": "LGORT"},
                {"as": "Area", "source": "LGOBE"},
            ],
            "alerts": {"enabled": True, "min_field": "stock_minimo", "qty_field": "StockLibre", "total_field": "ValorTotal"},
        }],
    }
    cfg_path.write_text(json.dumps(cfg))
    assert cli.cmd_migrate(str(cfg_path)) == 0

    # Crear tabla derivada con datos simulados
    eng = create_engine(f"sqlite:///{db_path}")
    df_sample = pd.DataFrame([
        {"ID": "MAT-001", "Descripcion": "Tornillo Acero", "StockLibre": 5, "stock_minimo": 10, "ValorTotal": 50000.0, "Centro": "C01", "Almacen": "A01", "Area": "Taller"},
        {"ID": "MAT-002", "Descripcion": "Tuerca 1/2", "StockLibre": 20, "stock_minimo": 15, "ValorTotal": 10000.0, "Centro": "C01", "Almacen": "A01", "Area": "Taller"},
    ])
    df_sample.to_sql("inv_bodega", eng, if_exists="replace", index=False)

    monkeypatch.setenv("ETL_CONFIG", str(cfg_path))
    monkeypatch.setenv("ETL_SECRET", "super-secret-key-32-chars-for-export-tests!")
    monkeypatch.setattr(auth_engine, "_ENGINE_PROVIDER", dash_app.get_engine)
    monkeypatch.setattr("dashboard.app.get_manager", lambda cfg: _licensed_manager())

    auth_users.create_user("admin", "adminpass", role="admin", is_root=True, must_change_password=False)
    auth_users.create_user("user", "userpass", role="user", is_root=False, must_change_password=False)
    return str(cfg_path)


def test_derived_export_excel(export_env):
    client = TestClient(dash_app.app)
    client.post("/api/auth/login", json={"username": "admin", "password": "adminpass"})

    res = client.get("/api/derived/inv_bodega/export/excel")
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.headers["content-type"]
    assert 'filename="inv_bodega.xlsx"' in res.headers["content-disposition"]

    # Verificar que openpyxl puede abrir el archivo y leer hojas
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Stock" in wb.sheetnames
    assert "StockBajo" in wb.sheetnames


def test_derived_export_csv(export_env):
    client = TestClient(dash_app.app)
    client.post("/api/auth/login", json={"username": "user", "password": "userpass"})

    res = client.get("/api/derived/inv_bodega/export/csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert 'filename="inv_bodega.csv"' in res.headers["content-disposition"]

    csv_text = res.content.decode("utf-8")
    assert "MAT-001" in csv_text
    assert "Tornillo Acero" in csv_text


def test_bi_export_success(export_env):
    client = TestClient(dash_app.app)
    client.post("/api/auth/login", json={"username": "admin", "password": "adminpass"})

    res = client.post("/api/bi/export")
    assert res.status_code == 200
    data = res.json()
    assert data["ok"] is True
    assert len(data["files"]) >= 1
    assert data["manifest"]["views"][0]["name"] == "inv_bodega"


def test_bi_manifest_and_guide(export_env):
    client = TestClient(dash_app.app)
    client.post("/api/auth/login", json={"username": "admin", "password": "adminpass"})

    res_manifest = client.get("/api/bi/manifest")
    assert res_manifest.status_code == 200
    assert "views" in res_manifest.json()

    res_guide = client.get("/api/bi/guide")
    assert res_guide.status_code == 200
    assert "Power BI" in res_guide.json()["guide_markdown"]


def test_bi_download_zip(export_env):
    client = TestClient(dash_app.app)
    client.post("/api/auth/login", json={"username": "admin", "password": "adminpass"})

    res = client.get("/api/bi/download-zip")
    assert res.status_code == 200
    assert "application/zip" in res.headers["content-type"]
    assert 'filename="novus_power_bi_dataset.zip"' in res.headers["content-disposition"]

    # Verificar contenido del ZIP
    with zipfile.ZipFile(io.BytesIO(res.content)) as z:
        names = z.namelist()
        assert "manifest.json" in names
        assert "GUIA_CONEXION_POWER_BI.md" in names
        assert any(n.endswith(".csv") for n in names)


def test_bi_download_individual_files(export_env):
    client = TestClient(dash_app.app)
    client.post("/api/auth/login", json={"username": "admin", "password": "adminpass"})

    res_p = client.get("/api/bi/download/inv_bodega/parquet")
    assert res_p.status_code in (200, 501)
    if res_p.status_code == 200:
        assert 'filename="inv_bodega.parquet"' in res_p.headers["content-disposition"]

    res_c = client.get("/api/bi/download/inv_bodega/csv")
    assert res_c.status_code == 200
    assert 'filename="inv_bodega.csv"' in res_c.headers["content-disposition"]

    res_m = client.get("/api/bi/download/manifest")
    assert res_m.status_code == 200
    assert 'filename="manifest.json"' in res_m.headers["content-disposition"]

    res_g = client.get("/api/bi/download/guide")
    assert res_g.status_code == 200
    assert 'filename="GUIA_CONEXION_POWER_BI.md"' in res_g.headers["content-disposition"]


def test_bi_export_rejected_without_bi_module(export_env, monkeypatch):
    monkeypatch.setattr("dashboard.app.get_manager", lambda cfg: _licensed_manager(modules={"dashboard": True, "bi": False}))
    client = TestClient(dash_app.app)
    client.post("/api/auth/login", json={"username": "admin", "password": "adminpass"})

    res = client.post("/api/bi/export")
    assert res.status_code == 403
    assert res.json()["detail"] == "modulo_no_contratado:bi"
