import json

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

import cli
from db.sinks import get_sink
from derived import build_inventory, run_derived
from derived.excel_export import export_report_from_table

MARA = pd.DataFrame([
    {"MATNR": "M1", "MTART": "ROH", "MATKL": "G1", "XCHPF": "", "MEINS": "UN"},
    {"MATNR": "M2", "MTART": "VERP", "MATKL": "G2", "XCHPF": "", "MEINS": "UN"},
    {"MATNR": "M3", "MTART": "HAWA", "MATKL": "G3", "XCHPF": "", "MEINS": "UN"},
    {"MATNR": "M4", "MTART": "ROH", "MATKL": "G1", "XCHPF": "", "MEINS": "UN"},
])

MARD = pd.DataFrame([
    {"MANDT": 100, "MATNR": "M1", "WERKS": "P1", "LGORT": "W1", "LABST": 10.0, "LMINB": 5.0},
    {"MANDT": 100, "MATNR": "M1", "WERKS": "P1", "LGORT": "W2", "LABST": 2.0, "LMINB": 3.0},
    {"MANDT": 100, "MATNR": "M2", "WERKS": "P1", "LGORT": "W1", "LABST": 7.0, "LMINB": 0.0},
    {"MANDT": 100, "MATNR": "M3", "WERKS": "P2", "LGORT": "W3", "LABST": 50.0, "LMINB": 10.0},
    {"MANDT": 100, "MATNR": "M1", "WERKS": "P2", "LGORT": "W1", "LABST": 1.0, "LMINB": None},
])

MBEW = pd.DataFrame([
    {"MATNR": "M1", "BWKEY": "P1", "LBKUM": 10.0, "SALK3": 100.0, "STPRS": 10.0, "VERPR": 9.5},
    {"MATNR": "M1", "BWKEY": "P2", "LBKUM": 1.0, "SALK3": 200.0, "STPRS": 20.0, "VERPR": 19.0},
    {"MATNR": "M2", "BWKEY": "P1", "LBKUM": 7.0, "SALK3": 70.0, "STPRS": 7.0, "VERPR": 6.5},
    {"MATNR": "M4", "BWKEY": "P1", "LBKUM": 9.0, "SALK3": 500.0, "STPRS": 50.0, "VERPR": 49.0},
])

MAKT = pd.DataFrame([
    {"MANDT": 100, "MATNR": "M1", "SPRAS": "ES", "MAKTX": "Desc Uno", "MAKTG": "Desc Uno"},
    {"MANDT": 100, "MATNR": "M1", "SPRAS": "EN", "MAKTX": "Desc One", "MAKTG": "Desc One"},
    {"MANDT": 100, "MATNR": "M2", "SPRAS": "ES", "MAKTX": "Desc Dos", "MAKTG": "Desc Dos"},
    {"MANDT": 100, "MATNR": "M3", "SPRAS": "ES", "MAKTX": "Desc Tres", "MAKTG": "Desc Tres"},
])

T001L = pd.DataFrame([
    {"MANDT": 100, "WERKS": "P1", "LGORT": "W1", "LGOBE": "Bodega Central"},
    {"MANDT": 100, "WERKS": "P1", "LGORT": "W2", "LGOBE": "Bodega Auxiliar"},
])

SOURCE_TABLES = {"MARA": MARA, "MARD": MARD, "MBEW": MBEW, "MAKT": MAKT, "T001L": T001L}

DEFAULT_TABLE_NAMES = {
    "MARA": "Mara_Data",
    "MARD": "Mard_Data",
    "MBEW": "Mbew_Data",
    "MAKT": "Makt_Data",
    "T001L": "T001L_Data",
}


def make_source_config(path, table_names=None, registry=None):
    table_names = table_names or DEFAULT_TABLE_NAMES
    tables = []
    for source in table_names:
        keys = {
            "MARA": ["MATNR"],
            "MARD": ["MANDT", "MATNR", "WERKS", "LGORT"],
            "MBEW": ["MATNR", "BWKEY"],
            "MAKT": ["MANDT", "MATNR", "SPRAS"],
            "T001L": ["MANDT", "WERKS", "LGORT"],
        }[source]
        tables.append({"source": source, "target": table_names[source], "keys": keys})
    return {
        "environment": "qa",
        "source": {"type": "csv", "config": {"directory": str(path)}},
        "database": {"dialect": "sqlite", "database": str(path / "db.sqlite")},
        "tables": tables,
        "fields": {src: list(df.columns) for src, df in SOURCE_TABLES.items()},
        "derived": _base_derived(path),
    }


def _base_derived(path):
    return {
        "name": "inv_bodega",
        "text_lang": "ES",
        "dominio": {"filters": [{"field": "MTART", "in": ["ROH", "VERP"]}]},
        "precio": {"field": "STPRS"},
        "valor": {"field": "SALK3"},
        "area": {"mapping": {"M1": "Insumos", "M2": "Insumos"}},
        "stock_minimo": {"source_field": "LMINB"},
        "keys": ["MATNR", "WERKS", "LGORT"],
        "excel": {"path": str(path / "inventario_bodega.xlsx")},
    }


@pytest.fixture
def engine(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path / 'db.sqlite'}")
    with eng.begin() as conn:
        for name, df in SOURCE_TABLES.items():
            df.to_sql(DEFAULT_TABLE_NAMES[name], conn, index=False)
    return eng


@pytest.fixture
def config(engine, tmp_path):
    return make_source_config(tmp_path)


def test_build_inventory_rule_domain_default_columns(engine, config):
    df = build_inventory(engine, config)

    assert list(df.columns) == [
        "MATNR", "Descripcion", "Centro", "Almacen",
        "StockLibre", "Precio", "ValorTotal", "Area", "stock_minimo",
    ]
    assert len(df) == 4

    def row(material, centro, almacen):
        r = df[(df["MATNR"] == material) & (df["Centro"] == centro) & (df["Almacen"] == almacen)]
        assert len(r) == 1
        return r.iloc[0]

    r = row("M1", "P1", "W1")
    assert r[["Descripcion", "StockLibre", "Precio", "ValorTotal", "Area", "stock_minimo"]].tolist() == [
        "Desc Uno", 10.0, 10.0, 100.0, "Insumos", 5.0]
    assert row("M1", "P1", "W2")[["StockLibre", "stock_minimo"]].tolist() == [2.0, 3.0]
    assert row("M2", "P1", "W1")[["StockLibre", "Precio", "ValorTotal"]].tolist() == [7.0, 7.0, 70.0]

    p2 = row("M1", "P2", "W1")
    assert p2["StockLibre"] == 1.0
    assert p2["Precio"] == 20.0
    assert p2["ValorTotal"] == 200.0
    assert pd.isna(p2["stock_minimo"])


def test_build_inventory_text_lang_en(engine, config):
    config["derived"]["text_lang"] = "EN"
    df = build_inventory(engine, config)
    desc = df.loc[df["MATNR"] == "M1", "Descripcion"].unique().tolist()
    assert desc == ["Desc One"]


def test_build_inventory_mbew_aggregate(engine, config):
    config["derived"]["valor"]["by"] = "aggregate"
    df = build_inventory(engine, config)
    m1 = df[df["MATNR"] == "M1"]
    assert (m1["ValorTotal"] == 300.0).all()
    assert (m1["Precio"] == 10.0).all()


def test_build_inventory_registry_dominio(engine, config, tmp_path):
    with engine.begin() as conn:
        pd.DataFrame([
            {"material": "M1", "centro": "P1", "almacen": "W1"},
            {"material": "M1", "centro": "P1", "almacen": "W2"},
            {"material": "M9", "centro": "P9", "almacen": "W9"},
        ]).to_sql("material_centro_almacen", conn, index=False)

    config["derived"]["dominio"] = {"registry_table": "material_centro_almacen"}
    df = build_inventory(engine, config)

    assert len(df) == 3
    sin_stock = df[df["MATNR"] == "M9"].iloc[0]
    assert sin_stock["StockLibre"] == 0.0
    assert sin_stock["Precio"] == 0.0
    assert sin_stock["ValorTotal"] == 0.0
    assert sin_stock["Descripcion"] != sin_stock["Descripcion"]  # NaN
    assert sin_stock["Area"] == ""


def test_build_inventory_registry_with_filter(engine, config, tmp_path):
    with engine.begin() as conn:
        pd.DataFrame([
            {"material": "M1", "centro": "P1", "almacen": "W1"},
            {"material": "M9", "centro": "P9", "almacen": "W9"},
        ]).to_sql("material_centro_almacen", conn, index=False)

    config["derived"]["dominio"] = {
        "registry_table": "material_centro_almacen",
        "filters": [{"field": "MTART", "in": ["ROH", "VERP"]}],
    }
    df = build_inventory(engine, config)
    assert len(df) == 1
    assert df.iloc[0]["MATNR"] == "M1"


def test_build_inventory_area_reference_table(engine, config, tmp_path):
    with engine.begin() as conn:
        pd.DataFrame([
            {"material": "M1", "area_id": 1},
            {"material": "M2", "area_id": 2},
        ]).to_sql("material_area", conn, index=False)
        pd.DataFrame([
            {"id": 1, "nombre": "Insumos"},
            {"id": 2, "nombre": "Exportación"},
        ]).to_sql("area", conn, index=False)

    config["derived"]["area"] = {
        "reference_table": "material_area",
        "join": {"table": "area", "on_local": "area_id", "on_foreign": "id"},
        "columns": {"key": "material", "value": "nombre"},
    }
    df = build_inventory(engine, config)
    assert df.loc[df["MATNR"] == "M1", "Area"].unique().tolist() == ["Insumos"]
    assert df.loc[df["MATNR"] == "M2", "Area"].unique().tolist() == ["Exportación"]


def test_build_inventory_minimo_reference_table_composite(engine, config, tmp_path):
    with engine.begin() as conn:
        pd.DataFrame([
            {"material": "M1", "centro": "P1", "almacen": "W1", "stock_minimo": 12.0},
        ]).to_sql("stock_minimo", conn, index=False)

    config["derived"]["stock_minimo"] = {
        "reference_table": "stock_minimo",
        "key_columns": ["material", "centro", "almacen"],
        "value_column": "stock_minimo",
    }
    df = build_inventory(engine, config)
    p1w1 = df[(df["MATNR"] == "M1") & (df["Centro"] == "P1") & (df["Almacen"] == "W1")]
    assert p1w1["stock_minimo"].tolist() == [12.0]
    other = df[(df["MATNR"] == "M1") & (df["Centro"] == "P1") & (df["Almacen"] == "W2")]
    assert pd.isna(other["stock_minimo"]).all()


def test_build_inventory_custom_columns(engine, config):
    config["derived"]["columns"] = [
        {"as": "Material", "source": "MATNR"},
        {"as": "Stock", "source": "LABST"},
    ]
    df = build_inventory(engine, config)
    assert list(df.columns) == ["Material", "Stock"]
    assert len(df) == 4


def test_build_inventory_compute_column(engine, config):
    config["derived"]["precio"] = {"field": "VERPR"}
    config["derived"]["valor"] = {"field": "VERPR"}
    config["derived"]["columns"] = [
        {"as": "MATNR", "source": "MATNR"},
        {"as": "Centro", "source": "WERKS"},
        {"as": "Almacen", "source": "LGORT"},
        {"as": "Stock", "source": "LABST"},
        {"as": "Unitario", "source": "VERPR"},
        {"as": "Total", "compute": "LABST * VERPR"},
    ]
    df = build_inventory(engine, config)
    r = df[(df["MATNR"] == "M1") & (df["Centro"] == "P1") & (df["Almacen"] == "W1")].iloc[0]
    assert r["Total"] == 95.0
    m2 = df[df["MATNR"] == "M2"].iloc[0]
    assert m2["Total"] == 45.5


def test_export_excel_reporte_structure(engine, config, tmp_path):
    config["derived"]["precio"] = {"field": "VERPR"}
    config["derived"]["valor"] = {"field": "VERPR"}
    config["derived"]["columns"] = [
        {"as": "MATNR", "source": "MATNR"},
        {"as": "Descripcion", "source": "MAKTX"},
        {"as": "Centro", "source": "WERKS"},
        {"as": "Almacen", "source": "LGORT"},
        {"as": "AlmacenDesc", "source": "LGOBE"},
        {"as": "StockLibre", "source": "LABST"},
        {"as": "UMB", "source": "MEINS"},
        {"as": "Precio", "source": "VERPR"},
        {"as": "ValorTotal", "compute": "LABST * VERPR"},
        {"as": "Area", "source": "AREA"},
        {"as": "stock_minimo", "source": "STOCK_MINIMO"},
    ]
    config["derived"]["excel"].update({
        "sheet": "Reporte",
        "columns": [
            {"as": "Columna1", "source": "MATNR"},
            {"as": "DescripcionMaterial", "source": "Descripcion"},
            {"as": "Centro", "source": "Centro"},
            {"as": "NumeroAlmacen", "source": "Almacen"},
            {"as": "DescripcionAlmacen", "source": "AlmacenDesc"},
            {"as": "Stock", "source": "StockLibre"},
            {"as": "UMB", "source": "UMB"},
            {"as": "ValorUnitario", "source": "Precio"},
            {"as": "ValorizacionTotal", "source": "ValorTotal"},
        ],
    })
    n, path = run_derived(engine=engine, sink=get_sink(engine), config=config)
    assert n == 4

    wb = load_workbook(path)
    assert wb.sheetnames == ["Reporte", "StockBajo"]
    rep = wb["Reporte"]
    assert [c.value for c in rep[1]] == [
        "Columna1", "DescripcionMaterial", "Centro", "NumeroAlmacen",
        "DescripcionAlmacen", "Stock", "UMB", "ValorUnitario", "ValorizacionTotal",
    ]
    assert rep.max_row == 5
    assert rep["A2"].value == "M1"
    assert rep["B2"].value == "Desc Uno"
    assert rep["C2"].value == "P1"
    assert rep["D2"].value == "W1"
    assert rep["E2"].value == "Bodega Central"
    assert rep["F2"].value == 10.0
    assert rep["G2"].value == "UN"
    assert rep["H2"].value == 9.5
    assert abs(rep["I2"].value - 95.0) < 0.001

    bajo = wb["StockBajo"]
    assert bajo.cell(row=4, column=1).value == "Conteo de materiales con stock bajo"
    assert bajo.cell(row=4, column=2).value == 1
    assert abs(bajo.cell(row=5, column=2).value - 19.0) < 0.001


def test_build_inventory_missing_base_raises(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path / 'empty.sqlite'}")
    cfg = make_source_config(tmp_path)
    with pytest.raises(ValueError, match="Mard_Data"):
        build_inventory(eng, cfg)


def test_run_derived_materializes_table_and_excel(engine, config, tmp_path):
    n, path = run_derived(engine=engine, sink=get_sink(engine), config=config)

    assert n == 4
    assert path == str(tmp_path / "inventario_bodega.xlsx")

    with engine.connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM inv_bodega")).scalar() == 4
        info = {r[1]: r[5] for r in conn.execute(text("PRAGMA table_info(inv_bodega)"))}
        assert info["MATNR"] == 1 and info["Centro"] == 2 and info["Almacen"] == 3

    wb = load_workbook(path)
    assert wb.sheetnames == ["Stock", "StockBajo"]
    stock = wb["Stock"]
    headers = [c.value for c in stock[1]]
    assert headers == [
        "MATNR", "Descripcion", "Centro", "Almacen",
        "StockLibre", "Precio", "ValorTotal", "Area", "stock_minimo",
    ]
    assert stock.max_row == 5

    bajo = wb["StockBajo"]
    assert bajo.max_row == 5
    assert bajo["A1"].value == "MATNR"
    assert bajo["B1"].value == "Descripcion"
    assert bajo.cell(row=4, column=1).value == "Conteo de materiales con stock bajo"
    assert bajo.cell(row=4, column=2).value == 1
    assert bajo.cell(row=5, column=1).value == "Valor total en riesgo"
    assert bajo.cell(row=5, column=2).value == 100.0


def test_export_report_from_table(engine, config, tmp_path):
    run_derived(engine=engine, sink=get_sink(engine), config=config)

    config["derived"]["excel"]["path"] = str(tmp_path / "reporte.xlsx")
    path = export_report_from_table(engine, config)

    wb = load_workbook(path)
    assert wb["Stock"].max_row == 5
    assert wb.sheetnames == ["Stock", "StockBajo"]


def test_export_report_from_table_missing_table(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path / 'vacio.sqlite'}")
    cfg = make_source_config(tmp_path)
    with pytest.raises(ValueError, match="inv_bodega"):
        export_report_from_table(eng, cfg)


def test_cli_reporte_generates_excel(engine, config, tmp_path):
    run_derived(engine=engine, sink=get_sink(engine), config=config)

    cfg_path = tmp_path / "config.json"
    cfg_path.write_text(json.dumps(config))

    assert cli.cmd_reporte(str(cfg_path)) == 0
    assert (tmp_path / "inventario_bodega.xlsx").exists()


CONTROL_SQL = """
CREATE TABLE etl_execution (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    start_time TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    end_time TEXT NULL,
    status TEXT NOT NULL,
    message TEXT NULL
);
CREATE TABLE etl_progress (
    table_name TEXT PRIMARY KEY,
    rows_loaded INTEGER NULL,
    status TEXT NULL,
    updated_at TEXT NULL DEFAULT CURRENT_TIMESTAMP,
    last_delta_value VARCHAR(40) NULL
);
CREATE TABLE etl_execution_tables (
    run_id INTEGER NOT NULL,
    table_name TEXT NOT NULL,
    chunks_total INTEGER NULL,
    chunks_ok INTEGER NULL,
    rows_extracted INTEGER NULL,
    duration_s REAL NULL,
    status TEXT NULL,
    PRIMARY KEY (run_id, table_name)
);
"""


def test_import_stock_minimo_and_derived_usage(engine, config, tmp_path):
    from derived.minimos_import import import_stock_minimo

    xlsx = tmp_path / "minimos.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        pd.DataFrame({
            "CODIGO": [" m1 ", "M2", "NOPE"],
            "STOCK MÍNIMO": [5, None, 99],
            "AREA": [" Insumos ", "", None],
        }).to_excel(w, sheet_name="PLANTA MINIMOS 2026", index=False)

    n = import_stock_minimo(engine, get_sink(engine), str(xlsx), sheet="PLANTA MINIMOS 2026",
                            columns={"material": "CODIGO", "stock_minimo": "STOCK MÍNIMO",
                                     "area": "AREA"})
    assert n == 3
    with engine.connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM stock_minimo")).scalar() == 3
        assert float(conn.execute(
            text("SELECT stock_minimo FROM stock_minimo WHERE material='M1'")).scalar()) == 5.0
        assert conn.execute(
            text("SELECT area FROM stock_minimo WHERE material='M1'")).scalar() == "Insumos"

    config["derived"]["stock_minimo"] = {
        "reference_table": "stock_minimo",
        "key_columns": ["material"],
        "value_column": "stock_minimo",
    }
    config["derived"]["area"] = {
        "reference_table": "stock_minimo",
        "columns": {"key": "material", "value": "area"},
    }

    df = build_inventory(engine, config)
    m1 = df[df["MATNR"] == "M1"]
    assert (m1["stock_minimo"] == 5.0).all()
    assert (m1["Area"] == "Insumos").all()
    m2 = df[df["MATNR"] == "M2"]
    assert pd.isna(m2["stock_minimo"]).all()
    assert (m2["Area"] == "").all()


def test_cli_importar_minimos(engine, config, tmp_path):
    xlsx = tmp_path / "minimos.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        pd.DataFrame({
            "CODIGO": ["M1", "M2"],
            "STOCK MÍNIMO": [5, 7],
            "AREA": ["Insumos", ""],
        }).to_excel(w, sheet_name="PLANTA MINIMOS 2026", index=False)

    cfg_path = tmp_path / "config.json"
    cfg_path.write_text(json.dumps(config))

    assert cli.cmd_importar_minimos(str(cfg_path), str(xlsx), hoja="PLANTA MINIMOS 2026",
                                    col_material="CODIGO", col_stock="STOCK MÍNIMO",
                                    col_area="AREA") == 0
    with engine.connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM stock_minimo")).scalar() == 2
        assert float(conn.execute(
            text("SELECT stock_minimo FROM stock_minimo WHERE material='M2'")).scalar()) == 7.0


def test_import_stock_minimo_default_columns_generic(engine, tmp_path):
    from derived.minimos_import import import_stock_minimo

    xlsx = tmp_path / "minimos.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        pd.DataFrame({
            "material": ["A1", "A2"],
            "stock_minimo": [5, 7],
            "area": ["Insumos", "Otra"],
        }).to_excel(w, sheet_name="Datos", index=False)

    n = import_stock_minimo(engine, get_sink(engine), str(xlsx), sheet="Datos")
    assert n == 2
    with engine.connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM stock_minimo")).scalar() == 2


def test_import_stock_minimo_requires_sheet(engine, tmp_path):
    from derived.minimos_import import import_stock_minimo

    xlsx = tmp_path / "minimos.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        pd.DataFrame({"material": ["A1"], "stock_minimo": [1]}).to_excel(w, index=False)
    with pytest.raises(ValueError, match="hoja"):
        import_stock_minimo(engine, get_sink(engine), str(xlsx), sheet=None)


def test_runner_generates_derived_table_and_excel(tmp_path, monkeypatch):
    import etl_runner

    eng = create_engine(f"sqlite:///{tmp_path / 'db.sqlite'}")
    with eng.begin() as conn:
        for stmt in CONTROL_SQL.split(";"):
            if stmt.strip():
                conn.execute(text(stmt))

    monkeypatch.setattr(etl_runner, "acquire_lock", lambda: None)
    monkeypatch.setattr(etl_runner, "release_lock", lambda: None)
    monkeypatch.setattr(etl_runner, "update_state", lambda **kw: None)
    monkeypatch.setattr(etl_runner, "get_engine", lambda config=None: eng)
    monkeypatch.setenv("ETL_LIVE_FILE", str(tmp_path / "live.json"))

    class FakeConnector:
        type = "csv"

        def connect(self):
            pass

        def close(self):
            pass

        def ping(self):
            return True

        def extract(self, table, fields, filters=None, batch_size=30000, max_wa_chars=None):
            data = {
                "MARD": {"MANDT": [100, 100], "MATNR": ["M1", "M2"], "WERKS": ["P1", "P1"],
                         "LGORT": ["L1", "L1"], "LABST": [10.0, 20.0], "LMINB": [5.0, 5.0]},
                "MARA": {"MATNR": ["M1", "M2"], "MTART": ["ROH", "VERP"], "MATKL": ["G", "G"],
                         "XCHPF": ["", ""], "MEINS": ["UN", "UN"]},
                "MAKT": {"MANDT": [100, 100], "MATNR": ["M1", "M2"], "SPRAS": ["ES", "ES"],
                         "MAKTX": ["Uno", "Dos"], "MAKTG": ["Uno", "Dos"]},
                "MBEW": {"MATNR": ["M1", "M2"], "BWKEY": ["P1", "P1"], "LBKUM": [10.0, 20.0],
                         "SALK3": [100.0, 200.0], "STPRS": [10.0, 20.0], "VERPR": [9.0, 19.0]},
            }
            return pd.DataFrame(data[table])

    monkeypatch.setattr(etl_runner, "get_source", lambda sc: FakeConnector())

    config = {
        "environment": "qa",
        "source": {"type": "csv", "config": {"directory": str(tmp_path)}},
        "database": {"dialect": "sqlite", "database": str(tmp_path / "db.sqlite")},
        "tables": [
            {"source": "MARD", "target": "mard", "keys": ["MANDT", "MATNR", "WERKS", "LGORT"]},
            {"source": "MARA", "target": "mara", "keys": ["MATNR"]},
            {"source": "MAKT", "target": "makt", "keys": ["MANDT", "MATNR", "SPRAS"]},
            {"source": "MBEW", "target": "mbew", "keys": ["MATNR", "BWKEY"]},
        ],
        "fields": {
            "MARD": ["MANDT", "MATNR", "WERKS", "LGORT", "LABST", "LMINB"],
            "MARA": ["MATNR", "MTART", "MATKL", "XCHPF", "MEINS"],
            "MAKT": ["MANDT", "MATNR", "SPRAS", "MAKTX", "MAKTG"],
            "MBEW": ["MATNR", "BWKEY", "LBKUM", "SALK3", "STPRS", "VERPR"],
        },
        "derived": {
            "name": "inv_bodega",
            "dominio": {"filters": [{"field": "MTART", "in": ["ROH", "VERP"]}]},
            "stock_minimo": {"source_field": "LMINB"},
            "area": {"mapping": {"M1": "Insumos", "M2": "Insumos"}},
            "excel": {"path": str(tmp_path / "inventario_bodega.xlsx")},
        },
    }
    monkeypatch.setattr(etl_runner, "load_config", lambda path=None: config)

    etl_runner.run_etl_job()

    with eng.connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM inv_bodega")).scalar() == 2
        row = conn.execute(
            text("SELECT table_name, status FROM etl_execution_tables "
                 "WHERE table_name='inv_bodega'")).first()
        assert row.status == "ok"
    assert (tmp_path / "inventario_bodega.xlsx").exists()


def test_derived_list_with_two_views_materializes_both(engine, config, tmp_path):
    from derived import run_deriveds

    config["derived"] = [
        config["derived"],
        {
            "name": "stock_por_centro",
            "dominio": {"filters": [{"field": "MTART", "in": ["ROH"]}]},
            "stock_minimo": {"source_field": "LMINB"},
            "columns": [
                {"as": "MATNR", "source": "MATNR"},
                {"as": "Centro", "source": "WERKS"},
                {"as": "Almacen", "source": "LGORT"},
                {"as": "StockLibre", "source": "LABST"},
            ],
            "keys": ["MATNR", "WERKS", "LGORT"],
            "excel": {"path": str(tmp_path / "stock_por_centro.xlsx"),
                      "alerts": {"enabled": False}},
        },
    ]

    results = run_deriveds(engine=engine, sink=get_sink(engine), config=config)
    names = [r[0] for r in results]
    assert names == ["inv_bodega", "stock_por_centro"]

    with engine.connect() as conn:
        assert conn.execute(text("SELECT COUNT(*) FROM inv_bodega")).scalar() == 4
        assert conn.execute(text("SELECT COUNT(*) FROM stock_por_centro")).scalar() == 3
    assert (tmp_path / "inventario_bodega.xlsx").exists()
    assert (tmp_path / "stock_por_centro.xlsx").exists()


def test_derived_list_duplicate_names_rejected(engine, config, tmp_path):
    from derived.views import derived_views

    config["derived"] = [
        {"name": "inv_bodega", "excel": {"path": str(tmp_path / "a.xlsx")}},
        {"name": "inv_bodega", "excel": {"path": str(tmp_path / "b.xlsx")}},
    ]
    with pytest.raises(ValueError, match="duplicado"):
        derived_views(config)


def test_runner_registers_telemetry_per_view(tmp_path, monkeypatch):
    import etl_runner

    eng = create_engine(f"sqlite:///{tmp_path / 'db.sqlite'}")
    with eng.begin() as conn:
        for stmt in CONTROL_SQL.split(";"):
            if stmt.strip():
                conn.execute(text(stmt))

    monkeypatch.setattr(etl_runner, "acquire_lock", lambda: None)
    monkeypatch.setattr(etl_runner, "release_lock", lambda: None)
    monkeypatch.setattr(etl_runner, "update_state", lambda **kw: None)
    monkeypatch.setattr(etl_runner, "get_engine", lambda config=None: eng)
    monkeypatch.setenv("ETL_LIVE_FILE", str(tmp_path / "live.json"))

    class FakeConnector:
        type = "csv"

        def connect(self):
            pass

        def close(self):
            pass

        def ping(self):
            return True

        def extract(self, table, fields, filters=None, batch_size=30000, max_wa_chars=None):
            data = {
                "MARD": {"MANDT": [100, 100], "MATNR": ["M1", "M2"], "WERKS": ["P1", "P1"],
                         "LGORT": ["L1", "L1"], "LABST": [10.0, 20.0], "LMINB": [5.0, 5.0]},
                "MARA": {"MATNR": ["M1", "M2"], "MTART": ["ROH", "VERP"], "MATKL": ["G", "G"],
                         "XCHPF": ["", ""], "MEINS": ["UN", "UN"]},
                "MAKT": {"MANDT": [100, 100], "MATNR": ["M1", "M2"], "SPRAS": ["ES", "ES"],
                         "MAKTX": ["Uno", "Dos"], "MAKTG": ["Uno", "Dos"]},
                "MBEW": {"MATNR": ["M1", "M2"], "BWKEY": ["P1", "P1"], "LBKUM": [10.0, 20.0],
                         "SALK3": [100.0, 200.0], "STPRS": [10.0, 20.0], "VERPR": [9.0, 19.0]},
            }
            return pd.DataFrame(data[table])

    monkeypatch.setattr(etl_runner, "get_source", lambda sc: FakeConnector())

    config = {
        "environment": "qa",
        "source": {"type": "csv", "config": {"directory": str(tmp_path)}},
        "database": {"dialect": "sqlite", "database": str(tmp_path / "db.sqlite")},
        "tables": [
            {"source": "MARD", "target": "mard", "keys": ["MANDT", "MATNR", "WERKS", "LGORT"]},
            {"source": "MARA", "target": "mara", "keys": ["MATNR"]},
            {"source": "MAKT", "target": "makt", "keys": ["MANDT", "MATNR", "SPRAS"]},
            {"source": "MBEW", "target": "mbew", "keys": ["MATNR", "BWKEY"]},
        ],
        "fields": {
            "MARD": ["MANDT", "MATNR", "WERKS", "LGORT", "LABST", "LMINB"],
            "MARA": ["MATNR", "MTART", "MATKL", "XCHPF", "MEINS"],
            "MAKT": ["MANDT", "MATNR", "SPRAS", "MAKTX", "MAKTG"],
            "MBEW": ["MATNR", "BWKEY", "LBKUM", "SALK3", "STPRS", "VERPR"],
        },
        "derived": [
            {
                "name": "inv_bodega",
                "dominio": {"filters": [{"field": "MTART", "in": ["ROH", "VERP"]}]},
                "stock_minimo": {"source_field": "LMINB"},
                "area": {"mapping": {"M1": "Insumos", "M2": "Insumos"}},
                "excel": {"path": str(tmp_path / "inventario_bodega.xlsx")},
            },
            {
                "name": "resumen_stock",
                "dominio": {"filters": [{"field": "MTART", "in": ["ROH"]}]},
                "columns": [
                    {"as": "MATNR", "source": "MATNR"},
                    {"as": "StockLibre", "source": "LABST"},
                ],
                "keys": ["MATNR"],
                "excel": {"path": str(tmp_path / "resumen_stock.xlsx"),
                          "alerts": {"enabled": False}},
            },
        ],
    }
    monkeypatch.setattr(etl_runner, "load_config", lambda path=None: config)

    etl_runner.run_etl_job()

    with eng.connect() as conn:
        for name, count in (("inv_bodega", 2), ("resumen_stock", 1)):
            assert conn.execute(text(f"SELECT COUNT(*) FROM {name}")).scalar() == count
            row = conn.execute(
                text("SELECT status FROM etl_execution_tables WHERE table_name=:n"),
                {"n": name}).first()
            assert row.status == "ok"
    assert (tmp_path / "inventario_bodega.xlsx").exists()
    assert (tmp_path / "resumen_stock.xlsx").exists()


def test_cli_reporte_selecciona_vision_por_tabla(engine, config, tmp_path):
    from derived import run_deriveds

    config["derived"] = [
        config["derived"],
        {
            "name": "resumen_stock",
            "columns": [
                {"as": "MATNR", "source": "MATNR"},
                {"as": "Centro", "source": "WERKS"},
                {"as": "Almacen", "source": "LGORT"},
                {"as": "StockLibre", "source": "LABST"},
            ],
            "keys": ["MATNR", "WERKS", "LGORT"],
            "excel": {"path": str(tmp_path / "resumen_stock.xlsx"),
                      "alerts": {"enabled": False}},
        },
    ]
    run_deriveds(engine=engine, sink=get_sink(engine), config=config)

    cfg_path = tmp_path / "config.json"
    cfg_path.write_text(json.dumps(config))

    assert cli.cmd_reporte(str(cfg_path), tabla="resumen_stock") == 0
    assert (tmp_path / "resumen_stock.xlsx").exists()
    assert cli.cmd_reporte(str(cfg_path), tabla="inexistente") == 1

