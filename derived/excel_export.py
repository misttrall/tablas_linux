"""Exportación a Excel de la tabla derivada (inventario de bodega).

Hoja principal con el detalle (replica las columnas del dashboard de stock) y,
opcionalmente, una hoja de alertas con los materiales cuyo stock está bajo el
mínimo y sus totales.
"""

import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from derived.db import read_table, table_exists

_NUMERIC_COLUMNS = {
    "StockLibre", "Precio", "ValorTotal", "stock_minimo",
    "Stock", "ValorUnitario", "ValorizacionTotal",
}


def _low_stock_mask(df, min_field, qty_field):
    qty = pd.to_numeric(df[qty_field], errors="coerce")
    minimum = pd.to_numeric(df[min_field], errors="coerce")
    return minimum.notna() & (qty < minimum)


def _project_columns(df, columns_cfg):
    """Proyecta/renombra las columnas del detalle (estructura del reporte cliente)."""
    if not columns_cfg:
        return df
    out = {}
    for c in columns_cfg:
        if isinstance(c, str):
            out[c] = df[c]
        elif "compute" in c:
            out[c["as"]] = df.eval(c["compute"])
        else:
            out[c["as"]] = df[c["source"]]
    return pd.DataFrame(out)


def _coerce_numeric(df):
    """Convierte a numérico las columnas conocidas (vienen como texto del SAP)."""
    out = df.copy()
    for c in _NUMERIC_COLUMNS.intersection(out.columns):
        out[c] = pd.to_numeric(out[c], errors="coerce")
    return out


def _numeric_cols(df):
    return {i + 1 for i, c in enumerate(df.columns) if pd.api.types.is_numeric_dtype(df[c])}


def _write_sheet(df, writer, name):
    df.to_excel(writer, sheet_name=name, index=False)


def _postprocess(path, sheet, alerts_sheet, has_alerts, numeric):
    wb = load_workbook(path)
    names = [sheet] + ([alerts_sheet] if has_alerts else [])
    for name in names:
        ws = wb[name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.row == 1 and isinstance(cell.value, str):
                    widths[cell.column] = max(widths.get(cell.column, 0), len(cell.value))
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 50)
        for col in numeric.get(name, set()):
            for cell in ws[get_column_letter(col)]:
                if cell.row > 1:
                    cell.number_format = "#,##0.##"
    wb.save(path)


def export_inventory_excel(df, derived):
    """Escribe el inventario en un Excel y devuelve la ruta absoluta."""
    cfg = derived.get("excel") or {}
    path = os.path.abspath(cfg.get("path") or "output/inventario_bodega.xlsx")
    sheet = cfg.get("sheet", "Stock")
    alerts_sheet = cfg.get("alerts_sheet", "StockBajo")

    alerts_cfg = derived.get("alerts", {})
    enabled = alerts_cfg.get("enabled", True)
    min_field = alerts_cfg.get("min_field", "stock_minimo")
    qty_field = alerts_cfg.get("qty_field", "StockLibre")
    total_field = alerts_cfg.get("total_field", "ValorTotal")

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)

    main_df = _coerce_numeric(_project_columns(df, cfg.get("columns")))
    numeric = {sheet: _numeric_cols(main_df)}

    has_alerts = enabled and min_field in df.columns and qty_field in df.columns
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_sheet(main_df, writer, sheet)
        if has_alerts:
            low = _coerce_numeric(df[_low_stock_mask(df, min_field, qty_field)])
            numeric[alerts_sheet] = _numeric_cols(low)
            _write_sheet(low, writer, alerts_sheet)
            ws = writer.sheets[alerts_sheet]
            n = len(low)
            if total_field in low.columns:
                total = pd.to_numeric(low[total_field], errors="coerce").fillna(0).sum()
            else:
                total = 0
            row0 = len(low) + 3
            ws.cell(row=row0, column=1, value="Conteo de materiales con stock bajo")
            ws.cell(row=row0, column=2, value=int(n)).number_format = "#,##0"
            ws.cell(row=row0 + 1, column=1, value="Valor total en riesgo")
            ws.cell(row=row0 + 1, column=2, value=float(total)).number_format = "#,##0.##"
            for cell in (ws.cell(row=row0, column=1), ws.cell(row=row0 + 1, column=1)):
                cell.font = Font(bold=True)

    _postprocess(path, sheet, alerts_sheet, has_alerts, numeric)
    return path


def export_report_from_table(engine, config, derived=None):
    """Regenera el Excel leyendo la tabla derivada ya materializada (sin SAP)."""
    from derived.views import resolve_view

    view = resolve_view(config, derived)
    target = view.get("name", "inv_bodega")
    if not table_exists(engine, target):
        raise ValueError(f"La tabla derivada '{target}' no existe en la BD destino (¿corrió el ETL?)")
    df = read_table(engine, target)
    return export_inventory_excel(df, view)
